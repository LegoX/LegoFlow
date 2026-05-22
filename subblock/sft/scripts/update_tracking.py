#!/usr/bin/env python3
"""Append a row to artifacts/实验追踪表.xlsx after a successful training run.

Usage (called by train.sh):
    python scripts/update_tracking.py --block-dir /path/to/block
"""
import argparse
import fcntl
import json
import os
import statistics
import sys
import tempfile
from pathlib import Path

import openpyxl
import yaml


def load_config(block_dir: Path) -> dict:
    with open(block_dir / "config.yaml", encoding="utf-8") as f:
        config = yaml.safe_load(f)
    return config["runtime_info"]["input"]


def extract_score(record: dict) -> dict | None:
    """Read score from legacy IM or PangUML v2 meta_info.unique_info."""
    score = record.get("_score")
    if isinstance(score, dict):
        return score

    meta_info = record.get("meta_info")
    unique_info = meta_info.get("unique_info") if isinstance(meta_info, dict) else None
    score = unique_info.get("_score") if isinstance(unique_info, dict) else None
    return score if isinstance(score, dict) else None


def compute_im_stats(im_path: Path) -> dict:
    """Compute turn count, score stats from IM JSONL."""
    turns_list = []
    scores_list = []

    with open(im_path, encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            record = json.loads(line)
            msgs = record.get("messages", [])
            turns_list.append(len(msgs))
            score = extract_score(record)
            if score:
                cs = score.get("composite_score")
                if cs is not None:
                    scores_list.append(cs)

    def fmt(vals, precision=0):
        if not vals:
            return ""
        mx, mn, mean = max(vals), min(vals), statistics.mean(vals)
        if precision == 0:
            return f"Max: {int(mx)}\nMin: {int(mn)}\nMean: {int(mean)}"
        return f"Max: {round(mx, precision)}\nMin: {round(mn, precision)}\nMean: {round(mean, precision)}"

    return {
        "turns": fmt(turns_list),
        "scores": fmt(scores_list, 4),
        "count": len(turns_list),
    }


def compute_token_stats(lf_path: Path, model_path: str, trust_remote_code: bool) -> str:
    """Compute token length stats from LF JSON using the model tokenizer."""
    try:
        from transformers import AutoTokenizer

        tokenizer = AutoTokenizer.from_pretrained(
            model_path, trust_remote_code=trust_remote_code
        )
    except Exception:
        return ""

    with open(lf_path, encoding="utf-8") as f:
        records = json.load(f)

    lengths = []
    for rec in records:
        text = ""
        for msg in rec.get("messages", []):
            content = msg.get("content", "")
            if content:
                text += content
        token_ids = tokenizer.encode(text, add_special_tokens=False)
        lengths.append(len(token_ids))

    if not lengths:
        return ""
    return f"Max: {max(lengths)}\nMin: {min(lengths)}\nMean: {int(statistics.mean(lengths))}"


def resolve_output_dir(block_dir: Path, output_dir: str) -> Path:
    path = Path(output_dir)
    if path.is_absolute():
        return path
    return block_dir / "artifacts" / "model" / path.name


def derive_scaffold_label(scaffold: str, job_dir: str) -> str:
    """Try to extract scaffold + version from job_dir name."""
    name = Path(job_dir).name if job_dir else ""
    # e.g. swerebench-filtered-oraclesolved-openhands-sdk-1.14.0-GLM-5-...
    scaffold_map = {
        "openhands-sdk": "openhands-sdk",
        "claude-code": "claude-code",
        "open-code": "opencode",
        "terminus2": "terminus-2",
        "openhands": "openhands",
    }
    label = scaffold_map.get(scaffold, scaffold)
    # try to find version number after scaffold name in the directory
    import re
    for key in ["openhands-sdk", "claude-code", "opencode", "terminus-2", "openhands"]:
        pattern = rf"{re.escape(key)}-(\d+\.\d+(?:\.\d+)?)"
        m = re.search(pattern, name, re.IGNORECASE)
        if m:
            return f"{label} {m.group(1)}"
    return label


def derive_teacher_model(job_dir: str) -> str:
    """Extract teacher model name from job directory path."""
    import re
    name = Path(job_dir).name if job_dir else ""
    m = re.search(r"(GLM-\d+|GPT-\d+|Claude-\d+)", name, re.IGNORECASE)
    if m:
        return m.group(1).lower()
    return ""


def derive_think_mode(template: str, output_dir: str) -> str:
    basename = Path(output_dir).name if output_dir else ""
    if basename.endswith("_think"):
        return "think"
    if basename.endswith("_nothink"):
        return "nothink"
    if "nothink" in template:
        return "nothink"
    return "think"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--block-dir", required=True)
    parser.add_argument("--skip-token-stats", action="store_true",
                        help="Skip slow tokenizer-based token length computation")
    args = parser.parse_args()

    block = Path(args.block_dir)
    cfg = load_config(block)

    scaffold = cfg["source"]["scaffold"]
    job_dir = cfg["source"].get("job_dir", "") or ""
    data_name = cfg["conversion"]["data_name"]
    model_path = cfg["model"]["model_name_or_path"]
    trust_remote_code = bool(cfg["model"].get("trust_remote_code", False))
    template = cfg["training"]["template"]
    output_dir = cfg["training"]["output_dir"]

    output_basename = Path(output_dir).name
    abs_output_dir = resolve_output_dir(block, output_dir)

    im_path = block / "artifacts" / "data" / "im_data" / f"{data_name}.jsonl"
    lf_path = block / "artifacts" / "data" / "lf_data" / f"{data_name}.json"
    train_yaml = block / "artifacts" / "training_config" / f"{output_basename}.yaml"

    # Converter module command
    converter_map = {
        "openhands-sdk": "swe_data_process.openhands.convert_openhands_sdk_to_im",
        "claude-code": "swe_data_process.claudecode_opencode.convert_cc_to_im",
        "open-code": "swe_data_process.claudecode_opencode.convert_oc_to_im",
        "terminus2": "swe_data_process.terminus2.convert_terminus2_to_im",
    }
    converter_module = converter_map.get(scaffold, "")
    converter_command = f"python -m {converter_module}" if converter_module else ""

    # Compute stats from IM data
    im_stats = {"turns": "", "scores": "", "count": ""}
    if im_path.exists():
        try:
            im_stats = compute_im_stats(im_path)
        except Exception as e:
            print(f"WARNING: failed to compute IM stats: {e}", file=sys.stderr)

    # Compute token length stats (slow — uses tokenizer)
    token_stats = ""
    if not args.skip_token_stats and lf_path.exists():
        try:
            print("Computing token length stats (this may take a while)...")
            token_stats = compute_token_stats(lf_path, model_path, trust_remote_code)
        except Exception as e:
            print(f"WARNING: failed to compute token stats: {e}", file=sys.stderr)

    # Trajectory count from LF data
    traj_count = ""
    if lf_path.exists():
        try:
            with open(lf_path) as f:
                traj_count = str(len(json.load(f)))
        except Exception:
            traj_count = str(im_stats.get("count", ""))

    # Build row values (A-S)
    row = {
        "A": "python",
        "B": _derive_dataset_label(job_dir, data_name),
        "C": "harbor",
        "D": derive_scaffold_label(scaffold, job_dir),
        "E": derive_think_mode(template, output_dir),
        "F": derive_teacher_model(job_dir),
        "G": job_dir,
        "H": converter_command,
        "I": str(lf_path),
        "J": token_stats,
        "K": im_stats["turns"],
        "L": im_stats["scores"],
        "M": traj_count,
        "N": Path(model_path).name,
        "O": str(train_yaml),
        "P": output_basename,
        "Q": str(abs_output_dir),
        "R": "",
        "S": "",
    }

    # Append to Excel
    xlsx_path = block / "artifacts" / "实验追踪表.xlsx"
    if not xlsx_path.exists():
        print(f"ERROR: tracking table not found: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    lock_path = xlsx_path.with_suffix(xlsx_path.suffix + ".lock")
    with lock_path.open("w") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb["实验追踪"]
        next_row = ws.max_row + 1

        for col_letter, value in row.items():
            ws[f"{col_letter}{next_row}"] = value

        fd, tmp_name = tempfile.mkstemp(
            prefix=f".{xlsx_path.stem}.", suffix=xlsx_path.suffix, dir=xlsx_path.parent
        )
        os.close(fd)
        try:
            wb.save(tmp_name)
            os.replace(tmp_name, xlsx_path)
        finally:
            if os.path.exists(tmp_name):
                os.unlink(tmp_name)
    print(f"=== Updated tracking table: row {next_row} in {xlsx_path} ===")


def _derive_dataset_label(job_dir: str, data_name: str) -> str:
    """Extract dataset label from job_dir or data_name."""
    import re
    name = Path(job_dir).name if job_dir else data_name
    # e.g. swerebench-filtered-oraclesolved-openhands-sdk-...
    m = re.match(r"(swerebench[\w-]*oraclesolved|swerebench[\w-]*)", name)
    if m:
        return m.group(1)
    return data_name


if __name__ == "__main__":
    main()
